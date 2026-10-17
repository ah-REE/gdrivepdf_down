#!/usr/bin/env python3
"""editable_pptx.py - rebuild a view-only Drive deck as an editable .pptx.

How it works:
  1. bootstrap Google's viewer session (gdrive_down.viewer_bootstrap)
  2. fetch every slide render at W px (default 1600)
  3. fetch the per-slide text layer (viewer/presspage) and rebuild lines
     with gap-based word joining
  4. mask the baked-in text on each render, then lay real editable
     text boxes on top (position, size, and sampled colour preserved)

Usage:
  python tools/editable_pptx.py "https://docs.google.com/presentation/d/<ID>/edit" -o out.pptx
  python tools/editable_pptx.py <file id> -o out.pptx --width 2000

Requires: requests, Pillow, python-pptx.
"""

import argparse
import io
import json
import os
import random
import re
import string
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import gdrive_down as gd  # noqa: E402

from PIL import Image, ImageDraw  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR  # noqa: E402
from pptx.util import Pt  # noqa: E402

FONT_NAME = "Calibri"
WIDTH_DEFAULT = 1600
SIZE_FACTOR = 0.92
BOLD_HEIGHT = 18


def lum(c):
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def clean(text):
    out = "".join(ch for ch in text if not (0xE000 <= ord(ch) <= 0xF8FF))
    return " ".join(out.split())


def jparse(text):
    if text.startswith(")]}'"):
        text = text.split("\n", 1)[1]
    return json.loads(text)


def collect_runs(node, out):
    """Presspage leaves look like [ [y, x, h, w], "text" ]."""
    if isinstance(node, list):
        if (
            len(node) == 2
            and isinstance(node[0], list)
            and len(node[0]) == 4
            and isinstance(node[1], str)
        ):
            y, x, h, w = node[0]
            out.append({"x": x, "y": y, "h": h, "w": w, "t": node[1]})
            return
        for c in node:
            collect_runs(c, out)


def lines_of(runs):
    if not runs:
        return []
    lines = {}
    for r in runs:
        key = None
        for k in lines:
            if abs(k - r["y"]) <= max(4, r["h"] * 0.45):
                key = k
                break
        if key is None:
            key = r["y"]
        lines.setdefault(key, []).append(r)
    out = []
    for key in sorted(lines):
        parts = sorted(lines[key], key=lambda r: r["x"])
        text = ""
        prev = None
        for p in parts:
            if prev is not None:
                gap = p["x"] - (prev["x"] + prev["w"])
                if gap > max(1.2, p["h"] * 0.11) and not text.endswith(" "):
                    text += " "
            text += p["t"]
            prev = p
        out.append(
            {
                "x": min(p["x"] for p in parts),
                "y": round(key),
                "w": max(p["x"] + p["w"] for p in parts) - min(p["x"] for p in parts),
                "h": max(p["h"] for p in parts),
                "text": text,
            }
        )
    return out


def fetch_slide_text(session, cfg, prev, page):
    r = session.get(
        "%s/viewer/%s&page=%d" % (gd._viewer_base(), cfg["presspage"], page),
        headers={"Referer": prev},
        timeout=(15, 60),
    )
    if r.status_code != 200:
        return []
    runs = []
    collect_runs(jparse(r.text), runs)
    return lines_of(runs)


def fetch_render(session, cfg, prev, page, width, fssid):
    base = gd._viewer_base() + "/viewer/" + cfg["page"]
    url = ("%s&auditContext=forDisplay&page=%d&w=%d&webp=true&fssid=%s"
           % (base, page, width, fssid))
    for attempt in range(4):
        try:
            r = session.get(url, headers={"Referer": prev}, timeout=(15, 120))
        except gd.requests.RequestException:
            time.sleep(0.5 + 0.3 * attempt)
            continue
        if r.status_code == 200 and len(r.content) > 200:
            return r.content
        time.sleep(0.4 + 0.3 * attempt)
    return None


def build_deck(slides, renders, out_path):
    prs = Presentation()
    slide_w, slide_h = Pt(960), Pt(540)
    prs.slide_width, prs.slide_height = slide_w, slide_h
    blank = prs.slide_layouts[6]
    stats = {"boxes": 0, "chars": 0, "skipped": 0}
    for i, (lines, blob) in enumerate(zip(slides, renders)):
        im = Image.open(io.BytesIO(blob)).convert("RGB")
        scale = im.width / 960.0
        draw = ImageDraw.Draw(im)
        items = []
        for ln in lines:
            text = clean(ln["text"])
            if len(text) < 2 or not any(ch.isalnum() for ch in text):
                stats["skipped"] += 1
                continue
            x0 = min(max(0.0, ln["x"] * scale), im.width - 1.0)
            y0 = min(max(0.0, ln["y"] * scale), im.height - 1.0)
            x1 = min(max(0.0, (ln["x"] + ln["w"]) * scale), im.width - 1.0)
            y1 = min(max(0.0, (ln["y"] + ln["h"]) * scale), im.height - 1.0)
            if x1 - x0 < 2 or y1 - y0 < 2:
                stats["skipped"] += 1
                continue
            ring = []
            top = int(max(0, y0 - 4))
            bottom = int(min(im.height - 1, y1 + 4))
            left = int(max(0, x0 - 4))
            right = int(min(im.width - 1, x1 + 4))
            for xx in range(int(x0), int(x1) + 1, 3):
                ring.append(im.getpixel((xx, top)))
                ring.append(im.getpixel((xx, bottom)))
            for yy in range(int(y0), int(y1) + 1, 3):
                ring.append(im.getpixel((left, yy)))
                ring.append(im.getpixel((right, yy)))
            ring.sort(key=lum)
            bg = ring[len(ring) // 2] if ring else (255, 255, 255)
            inner = []
            for xx in range(int(x0), int(x1) + 1, 3):
                for yy in range(int(y0), int(y1) + 1, 3):
                    inner.append(im.getpixel((xx, yy)))
            cols = [c for c in inner if abs(lum(c) - lum(bg)) > 70]
            if cols:
                cols.sort(key=lum)
                tcol = cols[0] if lum(bg) > 127 else cols[-1]
            else:
                tcol = (0, 0, 0) if lum(bg) > 127 else (255, 255, 255)
            draw.rectangle([x0 - 2, y0 - 1, x1 + 2, y1 + 1], fill=bg)
            items.append((ln, text, tcol))
            stats["boxes"] += 1
            stats["chars"] += len(text)
        slide = prs.slides.add_slide(blank)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=90)
        buf.seek(0)
        slide.shapes.add_picture(buf, 0, 0, width=slide_w, height=slide_h)
        for ln, text, tcol in items:
            box = slide.shapes.add_textbox(
                Pt(ln["x"] - 2), Pt(ln["y"] - 3), Pt(ln["w"] + 16), Pt(ln["h"] + 6)
            )
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            run = tf.paragraphs[0].add_run()
            run.text = text
            f = run.font
            f.size = Pt(round(ln["h"] * SIZE_FACTOR, 1))
            f.name = FONT_NAME
            f.bold = ln["h"] >= BOLD_HEIGHT
            f.color.rgb = RGBColor(*tcol)
    prs.save(str(out_path))
    return stats


def main(argv=None):
    ap = argparse.ArgumentParser(
        prog="editable_pptx",
        description="Rebuild a view-only Google Slides deck as an editable .pptx "
                    "(render backgrounds with masked text + real text boxes).",
    )
    ap.add_argument("link", help="Docs/Drive link or file id")
    ap.add_argument("-o", "--out", default="editable.pptx", help="output .pptx path")
    ap.add_argument("--width", type=int, default=WIDTH_DEFAULT,
                    help="render width in px (default %d)" % WIDTH_DEFAULT)
    ap.add_argument("--keep-json", default=None,
                    help="also dump the extracted text model to this path")
    args = ap.parse_args(argv)

    fid, kind = gd.parse_file_id(args.link)
    if not fid:
        print("[!] Could not find a file id in %r" % args.link)
        return 1

    session = gd.new_session()
    cfg, prev, fssid = gd.viewer_bootstrap(session, fid, verbose=True)
    if "page" not in cfg or "presspage" not in cfg:
        print("[!] This file does not expose a page/text layer (is it a deck?)")
        return 1
    meta = jparse(
        session.get(gd._viewer_base() + "/viewer/" + cfg["meta"],
                    headers={"Referer": prev}, timeout=(15, 60)).text
    )
    pages = int(meta.get("pages") or 0)
    print("slides:", pages)

    renders, slides = [], []
    for page in range(pages):
        blob = fetch_render(session, cfg, prev, page, args.width, fssid)
        if blob is None:
            print("[!] slide %d render unavailable" % page)
            return 1
        renders.append(blob)
        slides.append(fetch_slide_text(session, cfg, prev, page))
        time.sleep(0.08)
        if (page + 1) % 10 == 0:
            sys.stdout.write("\r  fetched %d/%d" % (page + 1, pages))
            sys.stdout.flush()
    sys.stdout.write("\n")

    if args.keep_json:
        dump = Path(args.keep_json)
        dump.parent.mkdir(parents=True, exist_ok=True)
        dump.write_text(
            json.dumps({"slides": slides}, ensure_ascii=False), encoding="utf-8"
        )

    out = Path(args.out)
    if out.exists():
        out = Path(gd.unique_path(str(out)))
    stats = build_deck(slides, renders, out)
    print("built:", stats["boxes"], "text boxes |", stats["chars"], "chars |",
          stats["skipped"], "skipped")
    print("[OK] %s (%.2f MB)" % (out.resolve(), out.stat().st_size / 1e6))
    return 0


if __name__ == "__main__":
    sys.exit(main())
